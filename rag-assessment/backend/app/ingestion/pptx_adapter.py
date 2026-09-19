from pathlib import Path
from typing import List, Union
from pptx import Presentation

from app.ingestion.base import BaseAdapter, Document


class PptxAdapter(BaseAdapter):
    """Adapter for extracting text and metadata from PPTX presentations using python-pptx."""

    def extract(self, file_path: Union[str, Path]) -> List[Document]:
        """
        Extract text per slide from a PPTX file.

        Args:
            file_path: Path to the PPTX file.

        Returns:
            List[Document]: One Document per slide with 1-indexed slide number and metadata.

        Raises:
            FileNotFoundError: If the file does not exist.
            RuntimeError: If the PPTX file is corrupted or cannot be parsed.
        """
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"PPTX file not found: {file_path}")

        try:
            prs = Presentation(str(path))
            total_slides = len(prs.slides)
            documents: List[Document] = []

            for slide_idx, slide in enumerate(prs.slides):
                slide_number = slide_idx + 1
                slide_texts: List[str] = []

                # Extract text from shapes (text boxes, titles, placeholders)
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            p_text = paragraph.text.strip()
                            if p_text:
                                slide_texts.append(p_text)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            row_text = " | ".join(
                                cell.text.strip() for cell in row.cells if cell.text.strip()
                            )
                            if row_text:
                                slide_texts.append(row_text)

                # Extract speaker notes if available
                if slide.has_notes_slide and slide.notes_slide:
                    notes_frame = slide.notes_slide.notes_text_frame
                    if notes_frame:
                        notes_text = notes_frame.text.strip()
                        if notes_text:
                            slide_texts.append(f"[Notes: {notes_text}]")

                full_slide_text = "\n".join(slide_texts).strip()

                documents.append(
                    Document(
                        text=full_slide_text,
                        source=path.name,
                        page=slide_number,
                        metadata={
                            "filename": path.name,
                            "format": "pptx",
                            "page": slide_number,
                            "total_slides": total_slides,
                        },
                    )
                )

            return documents

        except Exception as e:
            raise RuntimeError(f"Failed to read or parse PPTX file '{path.name}': {e}") from e
